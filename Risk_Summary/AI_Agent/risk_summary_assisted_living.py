from pathlib import Path
from datetime import datetime
import pandas as pd
from openai import OpenAI
from docx import Document

# =========================
# CONFIG
# =========================

OPENAI_API_KEY = "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"

# Automatically detect the current user's home folder
HOME = Path.home()

# Input folder
INPUT_DIR = HOME / "OneDrive - Reconnect Community Health Services (1)" / "data" / "risk_id_programs" / "als"

# Output folder
OUTPUT_DIR = HOME / "OneDrive - Reconnect Community Health Services (1)" / "data" / "risk_id_programs" / "docs" / "als"

MODEL = "gpt-4.1"

SAVE_OUTPUT_TO_TXT = True
SAVE_OUTPUT_TO_DOCX = True

PROMPT = """
You are a strict, evidence-based risk identification engine.

The structured records are already provided in the input.
Analyze them immediately.

Do NOT ask for records.
Do NOT ask clarifying questions.
Do NOT request more information.
Do NOT explain your reasoning.
Return only the final findings.

CORE RULE
You must identify ONLY risks that are explicitly supported by the record text.
If the text does not directly support a risk, do not include it.
When in doubt, leave it out.
False positives are worse than missed borderline cases.

NON-NEGOTIABLE RULES
1. Use only the text provided in the records.
2. Do not infer, speculate, generalize, or hallucinate.
3. Do not fill gaps with assumptions.
4. Do not treat routine care notes as risks.
5. Reference clients by ClientID only.
6. Never use names.
7. A client may appear in more than one category only if each category is directly supported by the text.
8. If evidence is weak, indirect, or ambiguous, exclude it.
9. Omit any category with no supported findings.
10. If no supported findings exist at all, return exactly:
No supported risks identified from the provided records.

RISK CATEGORIES
Use exactly these headings:

Medical Emergencies / Health Changes
Medication / Treatment Concerns
Safety / Environmental Concerns
Behavioural / Cognitive / Mental Health Concerns
Service / Care Coordination Risks
Abuse / Neglect / Injury Concerns

CATEGORY DEFINITIONS

1. Medical Emergencies / Health Changes
Include only when the note explicitly describes a current or recent health concern, symptom, decline, injury, or medically relevant change.

Examples that DO qualify when directly stated:
- not feeling well
- weakness or feeling weak
- dizziness
- fainting
- swelling
- bruising
- pain
- delirium
- confusion due to health status
- mobility decline
- inability to get out of bed
- infection signs
- concern requiring doctor, nurse, or urgent follow-up
- client said they needed to call their doctor because they felt unwell

Important:
- “not feeling well,” “weak,” or “needs to call doctor” should be included here when explicitly stated.
- Pain may belong here if it reflects a health concern.
- Do not increase severity beyond what is written.

2. Medication / Treatment Concerns
Include only when there is an explicit medication or treatment problem.

Examples that DO qualify:
- missed medication
- refused medication
- did not take medication
- confusion about medication
- wrong timing or adherence issue
- essential treatment not completed
- medication management problem

Examples that DO NOT qualify:
- medication was merely mentioned
- Tylenol or another medication was used and worked as expected
- medication was listed without any issue
- pain relieved by medication with no adherence problem

Important:
- A note is NOT a medication concern just because medication is mentioned.
- Only flag this category when there is an explicit problem with medication or treatment adherence.

3. Safety / Environmental Concerns
Include only when the note explicitly describes unsafe environmental conditions.

Examples that DO qualify:
- expired food
- pests or rodents
- fire hazards
- blocked walkways
- unsanitary conditions
- unsafe living conditions
- environmental hazards affecting health or mobility

Do NOT include general housekeeping or clutter unless the note clearly states it is unsafe.

4. Behavioural / Cognitive / Mental Health Concerns
Include only when the note explicitly describes behaviour, cognition, or emotional state that creates concern.

Examples that DO qualify:
- agitation
- repeated distress
- paranoia
- behavioural escalation
- client repeatedly checking doors/windows in a concerning way
- confusion
- anxiety requiring significant reassurance
- refusal of care creating risk

Do NOT include mild frustration, isolated preference, or ordinary interpersonal difficulty unless a real concern is directly described.

5. Service / Discharge / Care Coordination Risks
Include only when the note explicitly describes a risk to continuity of care, discharge process, equipment setup, consent process, or service coordination.

Examples that DO qualify:
- discharge delayed
- waiting on equipment required for discharge or safe return home
- no response from substitute decision-maker or family causing service delay
- care cannot proceed due to missing approvals or coordination failure
- clear service gap or continuity risk directly stated

Examples that DO NOT qualify:
- generic service planning
- routine referrals
- ordinary follow-up
- general discussion of future supports
- vague planning comments
- normal file review activity

Important:
- Do NOT create a service/discharge risk unless the note explicitly mentions discharge, equipment delay, consent delay, service interruption, file reopening need, or continuity-of-care risk.
- Do NOT infer this category from routine case management language.

6. Abuse / Neglect / Injury Concerns
Include only when the note explicitly supports injury, unexplained bruising, possible neglect, or harm.
Do not speculate.

DECISION TEST
Before adding any finding, ask:
1. Is there a direct phrase in the note supporting this risk?
2. Would a human reviewer be able to point to the exact wording?
3. Is this category the most accurate category?
4. Am I classifying this based on what is written, not what I suspect?

If any answer is no, exclude it.

DEDUPLICATION RULE
If the same ClientID appears multiple times for the same issue:
- combine into one bullet under the most appropriate category
- use the strongest directly supported evidence only
- do not repeat weak or redundant details

OUTPUT FORMAT
Return only in this structure:

Medical Emergencies / Health Changes
- ClientID: [ID]
  Summary: [brief factual summary]
  Evidence: [exact phrase or tight paraphrase from the note]

Medication / Treatment Concerns
- ClientID: [ID]
  Summary: [brief factual summary]
  Evidence: [exact phrase or tight paraphrase from the note]

Safety / Environmental Concerns
- ClientID: [ID]
  Summary: [brief factual summary]
  Evidence: [exact phrase or tight paraphrase from the note]

Behavioural / Cognitive / Mental Health Concerns
- ClientID: [ID]
  Summary: [brief factual summary]
  Evidence: [exact phrase or tight paraphrase from the note]

Service / Care Coordination Risks
- ClientID: [ID]
  Summary: [brief factual summary]
  Evidence: [exact phrase or tight paraphrase from the note]

Abuse / Neglect / Injury Concerns
- ClientID: [ID]
  Summary: [brief factual summary]
  Evidence: [exact phrase or tight paraphrase from the note]

FINAL CHECK BEFORE RESPONDING
- Remove any item based on implication only.
- Remove any item where the category is not explicitly supported.
- Remove medication findings if medication is only mentioned but no medication problem exists.
- Remove service/discharge findings unless the note explicitly supports care coordination risk.
- Include “not feeling well,” weakness, or need to call doctor under Medical when directly stated.
- Omit empty categories.
- Output only the final findings.
""".strip()

# =========================
# HELPERS
# =========================

def get_most_recent_csv_file(folder: Path) -> Path:
    csv_files = list(folder.glob("*.csv"))
    if not csv_files:
        raise FileNotFoundError(f"No .csv files found in {folder}")

    # Uses file creation time on Windows
    return max(csv_files, key=lambda p: p.stat().st_ctime)


def read_csv_as_text(file_path: Path) -> str:
    df = pd.read_csv(file_path)

    if df.empty:
        return "No records found in the CSV."

    # Convert each row into a readable block for the model
    records = []
    for _, row in df.iterrows():
        parts = []
        for col, val in row.items():
            val_str = "" if pd.isna(val) else str(val).strip()
            parts.append(f"{col}: {val_str}")
        records.append("\n".join(parts))

    return "\n---\n".join(records)


def call_openai(file_text: str) -> str:
    client = OpenAI(api_key=OPENAI_API_KEY)

    response = client.responses.create(
        model=MODEL,
        instructions=PROMPT,
        input=file_text
    )

    return response.output_text.strip()


def save_output_as_docx(content: str, source_file: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = output_dir / f"{source_file.stem}_openai_summary_{timestamp}.docx"

    doc = Document()
    doc.add_heading("Risk Identification Summary", level=0)
    doc.add_paragraph(f"Source file: {source_file.name}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")

    # Simple formatting:
    # make known category lines into headings, everything else into paragraphs
    headings = {
        "Medical Emergencies / Health Changes",
        "Medication / Treatment Concerns",
        "Safety / Environmental Concerns",
        "Behavioural / Cognitive / Mental Health Concerns",
        "Service / Care Coordination Risks",
        "Abuse / Neglect / Injury Concerns",
    }

    for line in content.splitlines():
        stripped = line.strip()

        if not stripped:
            doc.add_paragraph("")
        elif stripped in headings:
            doc.add_heading(stripped, level=1)
        else:
            doc.add_paragraph(stripped)

    doc.save(output_file)
    return output_file


def save_output_as_txt(content: str, source_file: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = output_dir / f"{source_file.stem}_openai_summary_{timestamp}.txt"
    output_file.write_text(content, encoding="utf-8")
    return output_file


# =========================
# MAIN
# =========================

def main():
    newest_csv = get_most_recent_csv_file(INPUT_DIR)
    csv_text = read_csv_as_text(newest_csv)

    if not csv_text.strip():
        raise ValueError(f"The CSV is empty or unreadable: {newest_csv}")

    result = call_openai(csv_text)

    print(f"Newest CSV: {newest_csv}")
    print("\n===== OPENAI OUTPUT =====\n")
    print(result)

    if SAVE_OUTPUT_TO_DOCX:
        docx_path = save_output_as_docx(result, newest_csv, OUTPUT_DIR)
        print(f"\nSaved Word document to: {docx_path}")

    #if SAVE_OUTPUT_TO_TXT:
    #    txt_path = save_output_as_txt(result, newest_csv, OUTPUT_DIR)
    #   print(f"Saved text output to: {txt_path}")


if __name__ == "__main__":
    main()