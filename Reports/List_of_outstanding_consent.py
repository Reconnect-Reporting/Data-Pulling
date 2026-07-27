import os
from datetime import datetime
from pathlib import Path
from typing import Tuple

import pandas as pd


# =========================
#         DEFAULT PATHS
# =========================
user = os.environ.get("UserName") or os.environ.get("USERNAME") or os.getlogin()

DEFAULT_TREAT_DIR = Path(
    r"C:\Users\=\OneDrive - Reconnect Community Health Services (1)\data\Clean Data\Treat".replace("=", user)
)
DEFAULT_ALAYA_DIR = Path(
    r"C:\Users\=\OneDrive - Reconnect Community Health Services (1)\data\Clean Data\AlayaCare".replace("=", user)
)


# =========================
#        HELPERS
# =========================
def _downloads_dir(fallback: Path) -> Path:
    """Try to resolve the user's Downloads folder (Windows friendly)."""
    candidates = [
        Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Downloads",
        Path.home() / "Downloads",
        Path.home() / "OneDrive" / "Downloads",
    ]
    for p in candidates:
        if p.exists():
            return p
    return fallback  # fallback if none exists


def newest_csv(prefix: str, base: Path) -> Path:
    files = sorted(base.glob(f"{prefix}*.csv"), key=lambda f: f.stat().st_mtime, reverse=True)
    if not files:
        raise FileNotFoundError(f"No '{prefix}*.csv' found in {base}")
    return files[0]


def newest_any(prefix: str, base: Path, exts=(".csv", ".xlsx", ".xls")) -> Path:
    files = []
    for ext in exts:
        files.extend(base.glob(f"{prefix}*{ext}"))
    files = sorted(files, key=lambda f: f.stat().st_mtime, reverse=True)
    if not files:
        raise FileNotFoundError(f"No '{prefix}*{exts}' found in {base}")
    return files[0]


def read_any(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "cp1252", "latin1"):
            try:
                return pd.read_csv(path, encoding=enc, keep_default_na=False)
            except UnicodeDecodeError:
                continue
        raise UnicodeDecodeError("Could not decode CSV with utf-8-sig/cp1252/latin1.")
    return pd.read_excel(path, dtype=str)


def _pick_col(df: pd.DataFrame, candidates):
    cols = {c.lower().strip(): c for c in df.columns}
    for cand in candidates:
        key = cand.lower().strip()
        if key in cols:
            return cols[key]
    for k, orig in cols.items():
        if any(c.lower().strip() in k for c in candidates):
            return orig
    raise KeyError(f"None of {candidates} found in columns: {list(df.columns)}")


def _nonempty_series(s: pd.Series) -> pd.Series:
    return s.notna() & s.astype(str).str.strip().ne("")


def _choose_excel_engine() -> str:
    try:
        import openpyxl  # noqa: F401
        return "openpyxl"
    except Exception:
        try:
            import xlsxwriter  # noqa: F401
            return "xlsxwriter"
        except Exception:
            return "openpyxl"


def _excel_col_letter(n: int) -> str:
    """1-based column index -> Excel letters (1=A)."""
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


# =========================
#    CROSS-SYSTEM HELPERS
# =========================
def get_sfs_consent_treat_ids(alaya_dir: Path) -> set[str]:
    """
    Find all Treat IDs in ADT whose AlayaCare ClientID has a Consent Form in Form Report.
    These Treat IDs should be considered consent-complete for MHA too.
    """
    adt_path = newest_any("ADT with Zone and Reference", alaya_dir)
    print(f"Reading ADT for reverse cross-check: {adt_path}")
    adt = read_any(adt_path)

    client_col_a = _pick_col(adt, ["ClientID", "Client ID", "ID"])

    try:
        treat_id_col = _pick_col(adt, ["Treat ID", "TreatID", "Reference", "Reference ID", "Treat Ref ID"])
    except KeyError:
        print("Treat ID column not found in ADT; skipping reverse cross-check from SFS to MHA.")
        return set()

    form_path = newest_any("Form Report", alaya_dir)
    print(f"Reading Form Report for reverse cross-check: {form_path}")
    forms = read_any(form_path)

    formname_col = _pick_col(forms, ["FormName", "Form Name"])
    client_col_f = _pick_col(forms, ["ClientID", "Client ID", "ID"])

    consent_forms = forms[
        _nonempty_series(forms[formname_col])
        & forms[formname_col].astype(str).str.strip().eq("Consent Form")
        & _nonempty_series(forms[client_col_f])
    ].copy()

    sfs_consent_client_ids = set(consent_forms[client_col_f].astype(str).str.strip())

    adt = adt[_nonempty_series(adt[client_col_a]) & _nonempty_series(adt[treat_id_col])].copy()
    adt[client_col_a] = adt[client_col_a].astype(str).str.strip()
    adt[treat_id_col] = adt[treat_id_col].astype(str).str.strip()

    sfs_consent_treat_ids = set(
        adt.loc[adt[client_col_a].isin(sfs_consent_client_ids), treat_id_col]
        .astype(str)
        .str.strip()
    )
    sfs_consent_treat_ids = {x for x in sfs_consent_treat_ids if x != ""}

    print(f"SFS consent mapped to {len(sfs_consent_treat_ids)} Treat IDs for MHA removal")
    return sfs_consent_treat_ids


# =========================
#      CORE PIPELINE
# =========================
def build_missing_mha(treat_dir: Path, alaya_dir: Path) -> pd.DataFrame:
    census_path = newest_csv("rpt_Census", treat_dir)
    print(f"Reading Census: {census_path}")
    census = read_any(census_path)

    discharge_col = _pick_col(census, ["DischargeDate", "Discharge Date"])
    census_active = census[census[discharge_col].astype(str).str.strip().eq("")].copy()

    id_col = _pick_col(census_active, ["ID", "Client ID", "ClientID"])
    name_col = _pick_col(census_active, ["ClientName", "Client Name"])
    program_col = _pick_col(census_active, ["Program"])
    staff_col = _pick_col(census_active, ["Staff", "PrimaryWorker", "Primary Worker", "Staff Name", "Staff Name:"])

    try:
        admit_col = _pick_col(
            census_active,
            ["AdmitDate", "Admit Date", "AdmissionDate", "Admission Date", "Admission Date/Time"],
        )
    except KeyError:
        admit_col = None

    assess_path = newest_csv("Assessment", treat_dir)
    print(f"Reading Assessment: {assess_path}")
    assess = read_any(assess_path)

    tool_col = _pick_col(assess, ["TOOLNAME", "Tool Name"])
    assess_id_col = _pick_col(assess, ["ID", "Client ID", "ClientID"])
    assess_filtered = assess[
        _nonempty_series(assess[tool_col])
        & assess[tool_col].astype(str).str.strip().eq("Reconnect Consent Assessment Tool")
        & _nonempty_series(assess[assess_id_col])
    ].copy()

    active_ids = set(
        census_active.loc[_nonempty_series(census_active[id_col]), id_col].astype(str).str.strip()
    )
    treat_assess_ids = set(
        assess_filtered.loc[_nonempty_series(assess_filtered[assess_id_col]), assess_id_col].astype(str).str.strip()
    )

    # NEW: anyone who has consent in SFS should also count as completed for MHA
    sfs_consent_treat_ids = get_sfs_consent_treat_ids(alaya_dir)

    completed_ids_any_system = treat_assess_ids | sfs_consent_treat_ids
    missing_ids = active_ids - completed_ids_any_system

    keep_cols = [name_col, id_col, staff_col, program_col]
    rename_map = {name_col: "ClientName", id_col: "ID", staff_col: "Staff", program_col: "Program"}
    if admit_col is not None:
        keep_cols.append(admit_col)
        rename_map[admit_col] = "AdmitDate"

    missing_df_mha = (
        census_active[census_active[id_col].astype(str).str.strip().isin(missing_ids)]
        .loc[:, keep_cols]
        .rename(columns=rename_map)
        .sort_values(["Program", "Staff", "ClientName"], kind="stable")
        .reset_index(drop=True)
        .drop_duplicates()
    )

    try:
        progfin_path = newest_any("rpt_ProgStaffFinance", treat_dir)
        print(f"Reading ProgStaffFinance: {progfin_path}")
        pf = read_any(progfin_path)

        id_pf_col = _pick_col(pf, ["ID", "Client ID", "ClientID"])
        date_pf_col = _pick_col(pf, ["Interaction Start Date", "InteractionStartDate", "Interaction Date", "Date"])
        vtype_pf_col = _pick_col(pf, ["Visit Type", "VisitType", "Type"])

        allowed_types = {"FaceHome", "FaceComm", "FaceOffice", "Video"}

        pf = pf[_nonempty_series(pf[id_pf_col])].copy()
        pf[vtype_pf_col] = pf[vtype_pf_col].astype(str).str.strip()
        pf = pf[pf[vtype_pf_col].isin(allowed_types)].copy()

        pf[date_pf_col] = pd.to_datetime(pf[date_pf_col], errors="coerce")
        pf = pf.dropna(subset=[date_pf_col]).copy()

        if "AdmitDate" in missing_df_mha.columns:
            missing_df_mha["AdmitDate"] = pd.to_datetime(missing_df_mha["AdmitDate"], errors="coerce")

        missing_df_mha["ID"] = missing_df_mha["ID"].astype(str).str.strip()
        pf["_ID"] = pf[id_pf_col].astype(str).str.strip()

        pf = pf[pf["_ID"].isin(set(missing_df_mha["ID"]))].copy()

        pf_tmp = pf.merge(
            missing_df_mha[["ID", "AdmitDate"]].rename(columns={"ID": "_ID"}),
            on="_ID",
            how="left",
        )

        pf_tmp = pf_tmp.dropna(subset=[date_pf_col, "AdmitDate"]).copy()
        pf_tmp = pf_tmp[pf_tmp[date_pf_col] > pf_tmp["AdmitDate"]].copy()

        first_visits = (
            pf_tmp.groupby("_ID", as_index=False)[date_pf_col]
            .min()
            .rename(columns={date_pf_col: "First Visit Date", "_ID": "ID"})
        )

        missing_df_mha = missing_df_mha.merge(first_visits[["ID", "First Visit Date"]], on="ID", how="left")

    except FileNotFoundError:
        print("rpt_ProgStaffFinance not found; skipping First Visit join")

    return missing_df_mha


def build_missing_sfs(alaya_dir: Path, treat_dir: Path) -> pd.DataFrame:
    adt_path = newest_any("ADT with Zone and Reference", alaya_dir)
    print(f"Reading ADT: {adt_path}")
    adt = read_any(adt_path)

    status_col = _pick_col(adt, ["Status"])
    client_col_a = _pick_col(adt, ["ClientID", "Client ID", "ID"])
    name_col_a = _pick_col(adt, ["ClientFullName", "Client Name", "ClientName"])
    program_col_a = _pick_col(adt, ["ServiceCodeDept", "Service Code Department Name"])
    zone_col = _pick_col(adt, ["Zone"])
    servicecode_col = _pick_col(adt, ["ServiceCode", "Service Code"])

    try:
        treat_id_col = _pick_col(adt, ["Treat ID", "TreatID", "Reference", "Reference ID", "Treat Ref ID"])
    except KeyError:
        treat_id_col = None

    try:
        service_start_col = _pick_col(adt, ["ServiceStartDate", "Service Start Date", "StartDate", "Start Date"])
    except KeyError:
        service_start_col = None

    try:
        first_visit_col = _pick_col(adt, ["First Visit Date", "FirstVisitDate", "First Visit", "First Contact Date"])
    except KeyError:
        first_visit_col = None

    adt_active = adt[
        _nonempty_series(adt[status_col]) & adt[status_col].astype(str).str.strip().str.lower().eq("active")
    ].copy()
    adt_active = adt_active[_nonempty_series(adt_active[client_col_a])].copy()

    form_path = newest_any("Form Report", alaya_dir)
    print(f"Reading Form Report: {form_path}")
    forms = read_any(form_path)

    formname_col = _pick_col(forms, ["FormName", "Form Name"])
    client_col_f = _pick_col(forms, ["ClientID", "Client ID", "ID"])

    consent_forms = forms[
        _nonempty_series(forms[formname_col]) & forms[formname_col].astype(str).str.strip().eq("Consent Form")
    ].copy()
    consent_forms = consent_forms[_nonempty_series(consent_forms[client_col_f])].copy()

    adt_ids = set(adt_active[client_col_a].astype(str).str.strip())
    consent_ids = set(consent_forms[client_col_f].astype(str).str.strip())
    missing_ids_sfs = adt_ids - consent_ids

    keep_cols = [client_col_a, name_col_a, program_col_a, zone_col, status_col, servicecode_col]
    rename_map = {
        client_col_a: "ClientID",
        name_col_a: "ClientName",
        program_col_a: "ServiceCodeDept",
        zone_col: "Zone",
        status_col: "Status",
        servicecode_col: "ServiceCode",
    }

    if treat_id_col is not None:
        keep_cols.append(treat_id_col)
        rename_map[treat_id_col] = "TreatID"

    if service_start_col is not None:
        keep_cols.append(service_start_col)
        rename_map[service_start_col] = "ServiceStartDate"

    if first_visit_col is not None:
        keep_cols.append(first_visit_col)
        rename_map[first_visit_col] = "FirstVisitDate"

    missing_df_sfs = (
        adt_active[adt_active[client_col_a].astype(str).str.strip().isin(missing_ids_sfs)]
        .loc[:, keep_cols]
        .rename(columns=rename_map)
        .sort_values(["ServiceCodeDept", "Zone", "ClientName"], kind="stable")
        .reset_index(drop=True)
    )
    missing_df_sfs["Consent Status"] = "Missing"
    missing_df_sfs = missing_df_sfs.drop_duplicates()
    missing_df_sfs = missing_df_sfs[~missing_df_sfs["ClientName"].astype(str).str.contains("Test", na=False)].copy()

    # If missing in AlayaCare, check Treat Assessment by Treat ID
    if "TreatID" in missing_df_sfs.columns:
        assess_path = newest_csv("Assessment", treat_dir)
        print(f"Reading Treat Assessment for SFS cross-check: {assess_path}")
        assess = read_any(assess_path)

        tool_col = _pick_col(assess, ["TOOLNAME", "Tool Name"])
        assess_id_col = _pick_col(assess, ["ID", "Client ID", "ClientID"])

        assess_filtered = assess[
            _nonempty_series(assess[tool_col])
            & assess[tool_col].astype(str).str.strip().eq("Reconnect Consent Assessment Tool")
            & _nonempty_series(assess[assess_id_col])
        ].copy()

        treat_consent_ids = set(assess_filtered[assess_id_col].astype(str).str.strip())

        missing_df_sfs["TreatID"] = missing_df_sfs["TreatID"].astype(str).str.strip()
        has_treat_consent = missing_df_sfs["TreatID"].isin(treat_consent_ids)

        missing_df_sfs.loc[has_treat_consent, "Consent Status"] = "Completed in Treat"
        missing_df_sfs = missing_df_sfs[~has_treat_consent].copy()

    return missing_df_sfs


def export_to_excel(mha_df: pd.DataFrame, sfs_df: pd.DataFrame, out_path: Path) -> Path:
    """
    Export tables and write a highlighted label:
      - MHA: "Number of Unique Clients: {n}" in J5
      - SFS: "Number of Unique Clients: {n}" in J5
    Uniqueness is nunique() on ID / ClientID.
    Also highlight duplicate IDs/ClientIDs in yellow (conditional formatting).
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    engine = _choose_excel_engine()

    mha_unique = mha_df["ID"].nunique(dropna=True) if "ID" in mha_df.columns else 0
    sfs_unique = sfs_df["ClientID"].nunique(dropna=True) if "ClientID" in sfs_df.columns else 0

    label_mha = f"Number of Unique Clients: {mha_unique}"
    label_sfs = f"Number of Unique Clients: {sfs_unique}"

    with pd.ExcelWriter(out_path, engine=engine) as writer:
        mha_df.to_excel(writer, sheet_name="MHA", index=False)
        sfs_df.to_excel(writer, sheet_name="SFS", index=False)

        if engine == "xlsxwriter":
            ws_mha = writer.sheets["MHA"]
            ws_sfs = writer.sheets["SFS"]

            fmt = writer.book.add_format({"bold": True, "bg_color": "#FFF2CC", "border": 1})
            ws_mha.write(4, 9, label_mha, fmt)  # J5
            ws_sfs.write(4, 9, label_sfs, fmt)  # J5

            yellow = writer.book.add_format({"bg_color": "#FFFF00"})

            if "ID" in mha_df.columns:
                col_idx = mha_df.columns.get_loc("ID")
                last_row = len(mha_df) + 1
                ws_mha.conditional_format(
                    1,
                    col_idx,
                    last_row,
                    col_idx,
                    {
                        "type": "formula",
                        "criteria": f'=COUNTIF(${_excel_col_letter(col_idx+1)}:${_excel_col_letter(col_idx+1)},{_excel_col_letter(col_idx+1)}2)>1',
                        "format": yellow,
                    },
                )

            if "ClientID" in sfs_df.columns:
                col_idx = sfs_df.columns.get_loc("ClientID")
                last_row = len(sfs_df) + 1
                ws_sfs.conditional_format(
                    1,
                    col_idx,
                    last_row,
                    col_idx,
                    {
                        "type": "formula",
                        "criteria": f'=COUNTIF(${_excel_col_letter(col_idx+1)}:${_excel_col_letter(col_idx+1)},{_excel_col_letter(col_idx+1)}2)>1',
                        "format": yellow,
                    },
                )

        else:
            from openpyxl.styles import Font, PatternFill, Border, Side
            from openpyxl.formatting.rule import FormulaRule

            ws_mha = writer.sheets["MHA"]
            ws_sfs = writer.sheets["SFS"]

            fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
            font = Font(bold=True)
            border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            cell_mha = ws_mha.cell(row=5, column=10)
            cell_mha.value = label_mha
            cell_mha.fill = fill
            cell_mha.font = font
            cell_mha.border = border

            cell_sfs = ws_sfs.cell(row=5, column=10)
            cell_sfs.value = label_sfs
            cell_sfs.fill = fill
            cell_sfs.font = font
            cell_sfs.border = border

            yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

            if "ID" in mha_df.columns:
                col_num = mha_df.columns.get_loc("ID") + 1
                col_letter = _excel_col_letter(col_num)
                last_row = len(mha_df) + 1
                rule = FormulaRule(formula=[f'COUNTIF(${col_letter}:${col_letter},{col_letter}2)>1'], fill=yellow_fill)
                ws_mha.conditional_formatting.add(f"{col_letter}2:{col_letter}{last_row}", rule)

            if "ClientID" in sfs_df.columns:
                col_num = sfs_df.columns.get_loc("ClientID") + 1
                col_letter = _excel_col_letter(col_num)
                last_row = len(sfs_df) + 1
                rule = FormulaRule(formula=[f'COUNTIF(${col_letter}:${col_letter},{col_letter}2)>1'], fill=yellow_fill)
                ws_sfs.conditional_formatting.add(f"{col_letter}2:{col_letter}{last_row}", rule)

    print(f"Saved Excel: {out_path}")
    print(f"[EXPORT] {label_mha}")
    print(f"[EXPORT] {label_sfs}")
    return out_path


# =========================
#           MAIN
# =========================
def main(
    treat_dir: Path = DEFAULT_TREAT_DIR,
    alaya_dir: Path = DEFAULT_ALAYA_DIR,
    out_path: Path | None = None,
) -> Tuple[Path, pd.DataFrame, pd.DataFrame]:
    if out_path is None:
        downloads = _downloads_dir(fallback=treat_dir)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = downloads / f"Consent_Missing_MHA_SFS_{ts}.xlsx"

    # UPDATED: pass alaya_dir into build_missing_mha
    missing_df_mha = build_missing_mha(treat_dir, alaya_dir)

    remove_programs = [
        "CBT Adult Group",
        "JAM",
        "ACT Social Group",
        "CBTP",
        "Central Intake",
        "DBT 101 Skills Group",
        "FOCUS",
        "Behavioural Therapy",
        "Drama Therapy",
        "Mindfulness Group",
        "Mainstay",
        "Dog assisted Therapy Group",
        "Music Therapy",
        "Addiction Focused P-E Group",
        "WeTo Social Group",
        "DBT Groups",
        "CDBT Skills Group",
        "DBT",
        "SOTO Social Group",
        "Art Therapy",
    ]

    missing_df_mha = missing_df_mha[~(missing_df_mha["Program"].isin(remove_programs))].copy()
    missing_df_mha["AdmitDate"] = pd.to_datetime(missing_df_mha["AdmitDate"], errors="coerce")
    missing_df_mha = missing_df_mha[~(missing_df_mha["ClientName"].str.contains("Anonymous", na=False))]

    start_date = pd.Timestamp("2025-04-01")
    end_date = pd.Timestamp.today() - pd.Timedelta(days=30)

    missing_df_mha = missing_df_mha[
        (missing_df_mha["AdmitDate"] > start_date) & (missing_df_mha["AdmitDate"] < end_date)
    ].copy()

    missing_df_sfs = build_missing_sfs(alaya_dir, treat_dir)
    missing_df_sfs = missing_df_sfs[~(missing_df_sfs["ServiceCode"].astype(str).str.contains("YorkW", na=False))]

    if "ServiceStartDate" in missing_df_sfs.columns:
        missing_df_sfs["ServiceStartDate"] = pd.to_datetime(missing_df_sfs["ServiceStartDate"], errors="coerce")
    if "FirstVisitDate" in missing_df_sfs.columns:
        missing_df_sfs["FirstVisitDate"] = pd.to_datetime(missing_df_sfs["FirstVisitDate"], errors="coerce")

    if "ServiceStartDate" in missing_df_sfs.columns:
        missing_df_sfs = missing_df_sfs[
            (missing_df_sfs["ServiceStartDate"] > pd.Timestamp("2025-04-01"))
            & (missing_df_sfs["ServiceStartDate"] < (pd.Timestamp.today() - pd.Timedelta(days=30)))
        ].copy()

    if "FirstVisitDate" in missing_df_sfs.columns and "ServiceStartDate" in missing_df_sfs.columns:
        missing_df_sfs = missing_df_sfs[
            (missing_df_sfs["FirstVisitDate"].isna())
            | (missing_df_sfs["FirstVisitDate"] > missing_df_sfs["ServiceStartDate"])
        ].copy()

    cm_mask = (
        missing_df_sfs["ServiceCodeDept"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.contains("Case Management", case=False, na=False)
    )
    cm_clients = set(missing_df_sfs.loc[cm_mask, "ClientID"].astype(str).str.strip())
    missing_df_sfs = missing_df_sfs[
        (~missing_df_sfs["ClientID"].astype(str).str.strip().isin(cm_clients)) | cm_mask
    ].copy()

    if "Zone" in missing_df_sfs.columns:
        missing_df_sfs = missing_df_sfs.drop(columns=["Zone"])

    missing_df_sfs = missing_df_sfs.drop_duplicates()
    missing_df_sfs = missing_df_sfs[missing_df_sfs["ServiceCodeDept"].astype(str).str.strip() != ""]

    cm_sup_path = alaya_dir / "CM Supervisors.csv"
    if not cm_sup_path.exists():
        cm_sup_path = newest_any("CM Supervisors", alaya_dir)

    print(f"Reading CM Supervisors: {cm_sup_path}")
    cm_sup = read_any(cm_sup_path)

    cm_id_col = _pick_col(cm_sup, ["AlayaCare Client ID", "Alayacare Client ID", "AlayaCareID", "AlayaCare ID"])
    cm_sup[cm_id_col] = cm_sup[cm_id_col].astype(str).str.strip()
    missing_df_sfs["ClientID"] = missing_df_sfs["ClientID"].astype(str).str.strip()

    missing_df_sfs = missing_df_sfs.merge(
        cm_sup,
        how="left",
        left_on="ClientID",
        right_on=cm_id_col,
    )

    out_path = export_to_excel(missing_df_mha, missing_df_sfs, out_path)
    return out_path, missing_df_mha, missing_df_sfs


if __name__ == "__main__":
    main()